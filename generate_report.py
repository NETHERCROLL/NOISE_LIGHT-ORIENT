from docx import Document
from docx.shared import Inches, Pt, RGBColor
from datetime import datetime

def build_docx_report(db, freq, surface, material_name, watts, wh_day, savings_fcfa, co2_kg, status_text):
    doc = Document()

    # En-tête
    title = doc.add_heading("NOISE LIGHT — Audit d'Impact Acoustique & ESG", level=0)
    title.runs[0].font.color.rgb = RGBColor(13, 110, 253)
    
    doc.add_paragraph(f"Date de certification : {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph("--------------------------------------------------------------------------------")

    # 1. Synthèse Énergétique
    doc.add_heading("1. Diagnostic Acoustique & Énergétique", level=1)
    p1 = doc.add_paragraph()
    p1.add_run(f"• Niveau sonore ambiant : ").bold = True
    p1.add_run(f"{db} dBA\n")
    p1.add_run(f"• Fréquence dominante : ").bold = True
    p1.add_run(f"{freq} Hz\n")
    p1.add_run(f"• Surface des capteurs : ").bold = True
    p1.add_run(f"{surface} m²\n")
    p1.add_run(f"• Matériau piézoélectrique : ").bold = True
    p1.add_run(f"{material_name}\n")
    p1.add_run(f"• Puissance instantanée générée : ").bold = True
    p1.add_run(f"{watts*1000:.2f} mW\n")
    p1.add_run(f"• Énergie journalière accumulée : ").bold = True
    p1.add_run(f"{wh_day:.1f} Wh/jour")

    # 2. Impact Financier & ESG
    doc.add_heading("2. Modélisation Financière & Empreinte ESG", level=1)
    p2 = doc.add_paragraph()
    p2.add_run(f"• Économies estimées : ").bold = True
    p2.add_run(f"{savings_fcfa:,.0f} FCFA / mois\n".replace(",", " "))
    p2.add_run(f"• Réduction d'empreinte carbone (CO2) : ").bold = True
    p2.add_run(f"{co2_kg:.2f} kg CO2 évités / mois")

    # 3. Maintenance Prédictive
    doc.add_heading("3. Rapport de Maintenance Prédictive", level=1)
    p3 = doc.add_paragraph()
    p3.add_run(f"Statut des équipements industriels : ").bold = True
    p3.add_run(f"{status_text}")

    # Pied de page
    doc.add_paragraph("\n--------------------------------------------------------------------------------")
    footer = doc.add_paragraph("Rapport certifié et généré automatiquement par la plateforme NOISE LIGHT (Python Engine).")
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.italic = True

    file_path = "Audit_NOISE_LIGHT.docx"
    doc.save(file_path)
    return file_path